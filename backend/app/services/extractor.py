from pathlib import Path
from typing import Any, Optional

import numpy as np
from docx import Document
from flask import current_app


_reader: Optional[Any] = None
_rapidocr_engine: Optional[Any] = None


def _get_easyocr_reader() -> Any:
    """Initialize EasyOCR once and reuse it across requests."""
    global _reader
    if _reader is None:
        try:
            import easyocr  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ImportError("easyocr is required for image resume extraction.") from exc
        _reader = easyocr.Reader(["en"], gpu=False)
    return _reader


def _get_rapidocr_engine() -> Any:
    """Initialize RapidOCR once and reuse it across requests."""
    global _rapidocr_engine
    if _rapidocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ImportError("rapidocr-onnxruntime is required for RapidOCR extraction.") from exc
        _rapidocr_engine = RapidOCR()
    return _rapidocr_engine


def _extract_with_easyocr(image_input: Any) -> list[str]:
    reader = _get_easyocr_reader()
    result = reader.readtext(image_input, detail=0, paragraph=True)
    return [line for line in result if isinstance(line, str) and line.strip()]


def _extract_with_rapidocr(image_input: Any) -> list[str]:
    engine = _get_rapidocr_engine()
    result, _elapsed = engine(image_input)
    if not result:
        return []

    lines: list[str] = []
    for item in result:
        # RapidOCR returns tuples in shape (box, text, score).
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            text = item[1]
            if isinstance(text, str) and text.strip():
                lines.append(text)
    return lines


def _extract_ocr_text(image_input: Any) -> str:
    """Extract text using configured engine with fallback for resilience."""
    engine = "rapidocr"
    try:
        engine = str(current_app.config.get("OCR_ENGINE", "rapidocr")).strip().lower()
    except RuntimeError:
        # Allows use in scripts/tests outside an app context.
        pass

    if engine == "easyocr":
        return "\n".join(_extract_with_easyocr(image_input)).strip()

    try:
        return "\n".join(_extract_with_rapidocr(image_input)).strip()
    except Exception:
        # Fallback keeps OCR available if RapidOCR model/runtime fails.
        return "\n".join(_extract_with_easyocr(image_input)).strip()


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text from PDF; fallback to OCR only when no embedded text exists."""
    try:
        import fitz
    except ImportError as exc:
        raise ImportError("PyMuPDF is required for PDF extraction.") from exc

    text_chunks = []
    pages_without_text = []
    with fitz.open(file_path) as doc:
        for index, page in enumerate(doc):
            page_text = str(page.get_text("text") or "")
            if page_text.strip():
                text_chunks.append(page_text)
            else:
                pages_without_text.append(index)

    direct_text = "\n".join(text_chunks).strip()
    if direct_text:
        return direct_text

    if not pages_without_text:
        return ""

    # OCR is expensive; perform it only for fully scanned/image-only PDFs.
    ocr_chunks = []
    with fitz.open(file_path) as doc:
        for page_index in pages_without_text:
            page = doc.load_page(page_index)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            page_array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            page_text = _extract_ocr_text(page_array)
            if page_text:
                ocr_chunks.append(page_text)

    return "\n".join(ocr_chunks).strip()


def extract_text_from_image(file_path: Path) -> str:
    """Extract text from image resumes using configured OCR engine."""
    return _extract_ocr_text(str(file_path))


def extract_resume_text(file_path: Path) -> str:
    """Route extraction based on file extension."""
    extension = file_path.suffix.lower().lstrip(".")
    if extension == "pdf":
        return extract_text_from_pdf(file_path)
    if extension in {"png", "jpg", "jpeg", "webp"}:
        return extract_text_from_image(file_path)
    if extension == "docx":
        return extract_text_from_docx(file_path)
    if extension == "txt":
        return extract_text_from_txt(file_path)
    raise ValueError("Unsupported file format. Supported: PDF, DOCX, TXT, PNG, JPG, JPEG, WEBP")


def extract_text_from_docx(file_path: Path) -> str:
    """Extract text from DOCX files."""
    try:
        doc = Document(file_path)
        text_chunks = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_chunks.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_chunks.append(cell.text)
        return "\n".join(text_chunks).strip()
    except Exception as exc:
        raise ValueError(f"Failed to extract DOCX: {exc}") from exc


def extract_text_from_txt(file_path: Path) -> str:
    """Extract text from plain TXT files."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read().strip()
    except UnicodeDecodeError:
        # Try different encoding
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                return f.read().strip()
        except Exception as exc:
            raise ValueError(f"Failed to extract TXT: {exc}") from exc
